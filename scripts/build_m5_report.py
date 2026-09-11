from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path("outputs/Football-Quant-V7-M5-2026-09-11.docx")


def shade(cell, color="DCE6F1"):
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(node)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    if widths:
        for col, width in zip(table.columns, widths, strict=False):
            col.width = Inches(width)
    for cell, text in zip(table.rows[0].cells, headers, strict=False):
        cell.text = text
        shade(cell)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row, strict=False):
            cell.text = str(text)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        borders.append(item)
    table._tbl.tblPr.append(borders)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.line_spacing = 1.1
    doc.add_paragraph()


def add_link(doc, label, url):
    paragraph = doc.add_paragraph()
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        for border in style.element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.22
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(12.5)


def section(doc, title, text):
    doc.add_heading(title, level=2)
    doc.add_paragraph(text)


def build():
    doc = Document()
    configure(doc)
    doc.add_heading("明早六点前足球赛前分析报告", 0)
    doc.add_paragraph("Football Quant V7 真实网页研究验收报告")
    doc.add_paragraph("生成时间：2026-09-11 19:15 北京时间")
    doc.add_paragraph("分析窗口：2026-09-11 18:49 至 2026-09-12 06:00 北京时间")
    doc.add_paragraph(
        "结论：本轮公开网页可确认多场赛程，但没有取得任何一场的完整、同一时点、可去水盘口。"
        "因此正式推荐为零；以下方向只能作为观察，不应按正式推荐下注。"
    )

    doc.add_heading("核心摘要", 1)
    add_table(
        doc,
        ("类别", "数量", "说明"),
        [
            ("公开发现并进入扫描", "6", "4场开赛时间已核验，2场时间未充分核验"),
            ("完整可定价盘口", "0", "无法计算可靠去水概率与EV"),
            ("正式推荐", "0", "全部PASS"),
            ("低置信观察方向", "4", "塞维利亚胜、佛罗伦萨不败、西汉姆胜倾向、哈茨胜倾向"),
            ("角球与罚牌候选", "0", "无完整盘口及独立历史计数数据"),
        ],
        (1.4, 0.65, 4.7),
    )
    doc.add_paragraph(
        "优先观察：塞维利亚胜。公开文章给出欧洲赔率2.10；原始隐含概率47.62%。"
        "因为缺少完整1X2对手盘、同口径历史输入和独立模型概率，去水市场概率、理论赔率、"
        "原始EV和风险调整EV均标记为缺失，评级PASS。"
    )

    doc.add_heading("全部赛事扫描表", 1)
    rows = [
        (
            "09-12 02:45",
            "斯滕豪斯穆尔 对 哈茨",
            "苏格兰联赛杯",
            "PASS",
            "缺完整盘口；杯赛轮换和小样本风险",
        ),
        (
            "09-12 03:00",
            "西汉姆联 对 雷克瑟姆",
            "英冠",
            "PASS",
            "缺完整赔率；双方三天内再赛，疲劳显著",
        ),
        ("09-12 03:00", "塞维利亚 对 瓦伦西亚", "西甲", "PASS", "仅取得单边2.10；无法去水和计算EV"),
        ("09-12 03:45", "威尼斯 对 佛罗伦萨", "意甲", "PASS", "缺完整盘口；双方均三连败且教练变动"),
        ("时间未核验", "雷恩 对 马赛", "法甲", "PASS", "搜索发现比赛，但开赛时间与完整盘口未核验"),
        (
            "时间未核验",
            "摩顿 对 利文斯顿",
            "苏格兰赛事",
            "PASS",
            "搜索发现比赛，但开赛时间与完整盘口未核验",
        ),
    ]
    add_table(
        doc,
        ("北京时间", "比赛", "赛事", "等级", "主要原因"),
        rows,
        (1.05, 1.95, 1.05, 0.62, 2.15),
    )

    doc.add_heading("观察方向", 1)
    add_table(
        doc,
        ("优先级", "比赛", "方向", "已知报价", "判断"),
        [
            ("1", "塞维利亚 对 瓦伦西亚", "塞维利亚胜", "2.10 欧赔", "低置信观察"),
            ("2", "威尼斯 对 佛罗伦萨", "佛罗伦萨不败", "缺失", "低置信观察"),
            ("3", "西汉姆联 对 雷克瑟姆", "西汉姆胜倾向", "缺失", "低置信观察"),
            ("4", "斯滕豪斯穆尔 对 哈茨", "哈茨胜倾向", "缺失", "低置信观察"),
        ],
        (0.6, 2.05, 1.4, 1.0, 1.35),
    )
    doc.add_paragraph("没有二串一或四串一；不能用缺失盘口的方向凑组合。")

    doc.add_heading("重点比赛分析", 1)
    doc.add_heading("塞维利亚 对 瓦伦西亚", 2)
    doc.add_paragraph("开赛：北京时间9月12日03:00。观察方向：塞维利亚胜。正式等级：PASS。")
    doc.add_paragraph(
        "支持面：塞维利亚前四轮拿到7分；瓦伦西亚只有1分、1个进球并丢9球。瓦伦西亚确认缺少"
        "Guido、Rioja、Diakhaby、Sadiq、Foulquier和Copete，防线甚至可能由Pepelu客串中卫。"
        "公开推荐文章给出塞维利亚胜2.10。"
    )
    doc.add_paragraph(
        "反对面：瓦伦西亚近四次作客该球场均有积分，主帅Corberán面对塞维利亚两胜两平。"
        "这是强烈的对位反证；塞维利亚并非稳定强队，不能只凭瓦伦西亚上轮0比5负于巴萨追热门。"
    )
    add_table(
        doc,
        ("指标", "数值", "状态"),
        [
            ("当前赔率", "2.10", "文章发布时赔率，可能变化"),
            ("原始隐含概率", "47.62%", "1 / 2.10"),
            ("去水市场概率", "缺失", "没有完整同一时点1X2"),
            ("模型概率", "缺失", "历史/xG输入不足，不生成虚假λ"),
            ("理论赔率", "缺失", "模型概率缺失"),
            ("原始EV", "缺失", "不能定价"),
            ("风险调整EV", "缺失", "不能定价"),
        ],
        (1.6, 1.5, 3.7),
    )

    doc.add_heading("西汉姆联 对 雷克瑟姆", 2)
    doc.add_paragraph("开赛：北京时间9月12日03:00。观察方向：西汉姆胜倾向。正式等级：PASS。")
    doc.add_paragraph(
        "西汉姆三连胜，周二曾在2比1落后时反胜博尔顿3比2，Bowen连续两场进球；雷克瑟姆近三场"
        "不败且只失一球，防守回升是真实反证。两队都在周二比赛，短休使轮换、下半场体能和临场"
        "节奏更难判断。没有可核验完整盘口，不能将“榜首+连胜”直接转成投注。"
    )

    doc.add_heading("威尼斯 对 佛罗伦萨", 2)
    doc.add_paragraph("开赛：北京时间9月12日03:45。观察方向：佛罗伦萨不败。正式等级：PASS。")
    doc.add_paragraph(
        "双方联赛前三轮均告负。威尼斯有四名中卫缺阵，防守组合不稳定；佛罗伦萨刚换帅，Vanoli"
        "上任仅数日，体系和首发均存在明显不确定性。客队前场可利用威尼斯中路缺员，但自身状态"
        "同样不足，因此不宜把纸面阵容优势放大为强推。"
    )

    doc.add_heading("斯滕豪斯穆尔 对 哈茨", 2)
    doc.add_paragraph("开赛：北京时间9月12日02:45。观察方向：哈茨胜倾向。正式等级：PASS。")
    doc.add_paragraph(
        "这是联赛杯淘汰赛。哈茨整体实力更强，公开分析倾向哈茨并看好较多进球；但斯滕豪斯穆尔"
        "此前1比0淘汰马瑟韦尔，杯赛战意和低级别球队的低位防守不应轻视。没有完整让球、大小球"
        "与轮换信息，强弱差不能单独构成推荐。"
    )

    doc.add_heading("角球和罚牌", 1)
    doc.add_paragraph(
        "本轮未取得角球大小/让球和罚牌大小/让分的完整公开盘口，也没有足够的双方主客场角球、"
        "黄牌与裁判同口径样本。Football Quant V7没有使用进球模型替代角球或罚牌模型，两个辅助"
        "市场全部PASS。"
    )

    doc.add_heading("数据缺失与覆盖限制", 1)
    for item in (
        "没有公开发现可交叉核验的完整同一时点1X2、亚洲盘或大小球。",
        "只有塞维利亚胜2.10为明确单边报价；单边赔率不能去水。",
        "雷恩对马赛、摩顿对利文斯顿由当日推荐文章发现，但开赛时间未被第二来源确认，因此不进入可执行池。",
        "公开搜索无法证明穷尽全球全部低级别赛事；扫描表代表本次实际发现并核验的集合。",
        "未取得可验证的8场同口径主客场历史、对手质量和xG输入，因此不生成λ和模型胜率。",
        "未取得角球、黄牌、裁判和完整辅助盘口，不能计算辅助市场概率或EV。",
        "官方首发尚未公布不构成自动扣分；已确认的瓦伦西亚伤停作为风险事实记录。",
    ):
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("数据来源和采集时间", 1)
    sources = [
        (
            "AS 塞维利亚对瓦伦西亚赛程与状态",
            "https://as.com/futbol/primera/sevilla-valencia-tv-a-que-hora-es-donde-y-como-ver-laliga-ea-sports-online-hoy-f202609-n/",
        ),
        (
            "AS 瓦伦西亚预计阵容和伤停",
            "https://as.com/futbol/primera/alineacion-posible-del-valencia-ante-el-sevilla-en-laliga-ea-sports-f202609-n/",
        ),
        (
            "talkSPORT 当日推荐及塞维利亚2.10",
            "https://talksport.com/football/4572485/talksport-tips-best-football-bets-friday-11th-september/",
        ),
        (
            "The Sun 西汉姆对雷克瑟姆近况",
            "https://www.thesun.co.uk/betting/40331536/west-ham-vs-wrexham-bet365-bonus-code/",
        ),
        (
            "The Sun 雷克瑟姆赛程开球时间",
            "https://www.thesun.ie/sport/17226632/wrexham-2026-27-championship-fixtures-sky-sports-kick-off/",
        ),
        (
            "Viola Nation 威尼斯对佛罗伦萨前瞻",
            "https://www.violanation.com/fiorentina-match-coverage/22881/venezia-vs-fiorentina-preview-serie-a-score-prediction",
        ),
        (
            "Scottish Sun 苏格兰联赛杯时间",
            "https://www.thescottishsun.co.uk/sport/16696888/rangers-celtic-premier-sports-cup-date-time/",
        ),
    ]
    doc.add_paragraph(
        "统一采集时间：2026-09-11 10:49至11:15 UTC。页面未提供盘口观察时间时留空；"
        "文章的发布时间不冒充盘口变动时间。"
    )
    for label, url in sources:
        add_link(doc, label, url)

    doc.add_heading("方法和风险提示", 1)
    doc.add_paragraph(
        "模型要求独立、有来源的历史输入和完整市场后才计算Poisson、Dixon-Coles、去水概率及EV。"
        "本次条件未满足，所以缺失值全部如实显示，未从赔率反推λ，也未把定性判断包装成模型概率。"
    )
    doc.add_paragraph(
        "PASS并不预测比赛不会按观察方向发展，只表示当前证据不足以形成可执行的量化推荐。"
        "报告仅供信息参考，不保证盈利；本项目不连接投注账户，不执行真实下注。"
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Football Quant V7  M5真实网页研究  2026-09-11"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
