"""Word rendering only: all conclusions come from the frozen Report."""

from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from football_quant.domain import Candidate, Grade, Market, Report

SHANGHAI = ZoneInfo("Asia/Shanghai")
ASIAN = {
    Market.HANDICAP,
    Market.TOTAL,
    Market.CORNER_TOTAL,
    Market.CORNER_HANDICAP,
    Market.CARD_TOTAL,
    Market.CARD_HANDICAP,
}
LABELS = {
    Market.RESULT: "胜平负",
    Market.HANDICAP: "亚洲让球",
    Market.TOTAL: "进球大小",
    Market.GOALS: "总进球数",
    Market.BTTS: "双方进球",
    Market.CORNER_TOTAL: "角球大小",
    Market.CORNER_HANDICAP: "角球让球",
    Market.CARD_TOTAL: "罚牌大小",
    Market.CARD_HANDICAP: "罚牌让分",
}
DIRECTIONS = {
    "home": "主队",
    "away": "客队",
    "draw": "平局",
    "over": "大",
    "under": "小",
    "yes": "是",
    "no": "否",
}


def percentage(value: float | None) -> str:
    return "缺失" if value is None else f"{value:.2%}"


def selection(candidate: Candidate) -> str:
    q = candidate.quote
    line = "" if q.line is None else f" {q.line:+}"
    display = DIRECTIONS.get(q.selection, q.selection)
    return f"{LABELS[q.market]} {display}{line}"


def display_odds(candidate: Candidate) -> str:
    q = candidate.quote
    if q.market in ASIAN:
        return f"香港水位 {q.decimal_odds - 1:.3f}"
    return f"欧赔 {q.decimal_odds:.3f}"


def configure(doc: DocumentType) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.8)
    section.top_margin = section.bottom_margin = Inches(0.7)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Noto Sans CJK SC"
        for border in style.element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(12)


def table(doc: DocumentType, headers: tuple[str, ...], rows: list[tuple[str, ...]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for cell, label in zip(t.rows[0].cells, headers, strict=True):
        cell.text = label
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "DCE6F1")
        cell._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, value in zip(t.add_row().cells, row, strict=True):
            cell.text = value
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            item.set(qn(f"w:{key}"), value)
        borders.append(item)
    t._tbl.tblPr.append(borders)
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment = 1
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
    doc.add_paragraph()


def hyperlink(doc: DocumentType, text: str, url: str) -> None:
    p = doc.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(
        qn("r:id"),
        p.part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        ),
    )
    run, content = OxmlElement("w:r"), OxmlElement("w:t")
    content.text = text
    run.append(content)
    link.append(run)
    p._p.append(link)


def summary(report: Report) -> str:
    counts = {grade: 0 for grade in Grade}
    for match in report.matches:
        for candidate in match.candidates:
            counts[candidate.grade] += 1
    return (
        f"扫描 {len(report.matches)} 场；候选等级："
        + "、".join(f"{grade.value} {counts[grade]}" for grade in Grade)
        + "。仅单场方向，不生成串关。"
    )


def candidate_detail(doc: DocumentType, candidate: Candidate) -> None:
    c, p = candidate, candidate.price
    doc.add_heading(f"{selection(c)}  {c.grade.value}", level=3)
    doc.add_paragraph(
        f"{display_odds(c)}；综合置信评分 {percentage(c.confidence)}"
        f"（非胜率）；完整度 {percentage(c.completeness)}。"
    )
    if p:
        table(
            doc,
            ("模型有效概率", "原始隐含概率", "去水市场概率"),
            [
                (
                    percentage(p.model_probability),
                    percentage(p.implied_probability),
                    percentage(p.devig_probability),
                )
            ],
        )
        fair = "不可定价" if p.fair_odds is None else f"{p.fair_odds:.3f}"
        doc.add_paragraph(
            f"理论欧赔 {fair}；当前欧赔 {p.current_odds:.3f}；"
            f"原始 EV {p.ev:+.2%}；风险调整 EV {p.risk_ev:+.2%}。"
        )
        doc.add_paragraph(
            "结算概率（全赢／半赢／走／半输／全输）：" + "／".join(percentage(v) for v in p.states)
        )
    for title, items in (
        ("支持理由", c.supports),
        ("反对理由", c.objections),
        ("风险", c.risks),
        ("PASS 原因", c.reasons),
    ):
        doc.add_paragraph(f"{title}：" + ("；".join(items) if items else "无额外记录"))


def write_report(report: Report, path: Path) -> None:
    doc = Document()
    configure(doc)
    title = "足球量化分析测试报告" if report.mode.value == "test" else "足球赛前量化分析报告"
    doc.add_heading(title, 0)
    doc.add_paragraph(
        "测试数据仅用于验证链路，不能用于投注。"
        if report.mode.value == "test"
        else "本报告分析本次公开发现的赛事，结论以已采集证据为依据。"
    )
    doc.add_paragraph(f"生成时间：{report.generated.astimezone(SHANGHAI):%Y-%m-%d %H:%M} 北京时间")
    doc.add_paragraph(
        f"分析窗口：{report.started.astimezone(SHANGHAI):%Y-%m-%d %H:%M} 至 "
        f"{report.deadline.astimezone(SHANGHAI):%Y-%m-%d %H:%M} 北京时间"
    )
    doc.add_heading("核心摘要与数据完整度", 1)
    doc.add_paragraph(summary(report))
    for note in report.coverage_notes:
        doc.add_paragraph(note)
    doc.add_heading("全部赛事扫描表", 1)
    rows = [
        (
            m.fixture.kickoff.astimezone(SHANGHAI).strftime("%m-%d %H:%M"),
            f"{m.fixture.home} 对 {m.fixture.away}",
            m.fixture.competition,
            "／".join(sorted({c.grade.value for c in m.candidates})) or "PASS",
        )
        for m in report.matches
    ]
    table(doc, ("北京时间", "比赛", "赛事", "等级"), rows)
    pools(doc, report)
    doc.add_heading("深度比赛分析与概率定价", 1)
    for match in report.matches:
        doc.add_heading(f"{match.fixture.home} 对 {match.fixture.away}", 2)
        for note in match.notes:
            doc.add_paragraph(note)
        doc.add_paragraph(f"主队预期进球 {match.lambda_home}；客队预期进球 {match.lambda_away}。")
        doc.add_paragraph("常见比分：" + "、".join(f"{i}:{j} {p:.1%}" for i, j, p in match.scores))
        for c in match.candidates:
            candidate_detail(doc, c)
    ending(doc, report)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def pools(doc: DocumentType, report: Report) -> None:
    groups = (
        (
            "主市场候选池",
            set(Market)
            - {
                Market.CORNER_TOTAL,
                Market.CORNER_HANDICAP,
                Market.CARD_TOTAL,
                Market.CARD_HANDICAP,
            },
        ),
        ("角球候选池", {Market.CORNER_TOTAL, Market.CORNER_HANDICAP}),
        ("罚牌候选池", {Market.CARD_TOTAL, Market.CARD_HANDICAP}),
    )
    for title, markets in groups:
        doc.add_heading(title, 1)
        rows = [
            (f"{m.fixture.home} 对 {m.fixture.away}", selection(c), c.grade.value)
            for m in report.matches
            for c in m.candidates
            if c.quote.market in markets
        ]
        if rows:
            table(doc, ("比赛", "方向", "等级"), rows)
        else:
            doc.add_paragraph("无候选：本次未取得该市场可定价数据。")
    doc.add_heading("低置信方向与 PASS", 1)
    for m in report.matches:
        if not m.candidates:
            doc.add_paragraph(
                f"{m.fixture.home} 对 {m.fixture.away}：PASS；" + "；".join(m.missing)
            )
        for c in m.candidates:
            if c.grade in (Grade.C, Grade.PASS):
                doc.add_paragraph(
                    f"{m.fixture.home} 对 {m.fixture.away} {selection(c)}："
                    f"{c.grade.value}；" + "；".join(c.reasons or c.risks)
                )


def ending(doc: DocumentType, report: Report) -> None:
    doc.add_heading("数据缺失与来源冲突", 1)
    for m in report.matches:
        doc.add_paragraph(
            f"{m.fixture.home} 对 {m.fixture.away}："
            + "；".join(m.missing + m.conflicts or ("未记录冲突",))
        )
    doc.add_heading("数据来源与采集时间", 1)
    for e in report.evidence:
        hyperlink(doc, f"{e.id} {e.source_name} — {e.source_url}", e.source_url)
        doc.add_paragraph(
            f"采集 UTC {e.retrieved_at_utc.isoformat()}；页面观察时间 "
            f"{e.observed_at_utc or '未提供'}；核验 {e.validation_status.value}"
        )
    doc.add_heading("方法与风险说明", 1)
    doc.add_paragraph(
        "比分采用 Poisson 和 Dixon–Coles；亚洲盘使用全赢、半赢、走、半输、全输"
        "逐状态定价。亚洲盘模型有效概率不是全赢概率。"
    )
    doc.add_paragraph(
        "EV 为每单位投入的模型收益估计，风险调整扣除明示不确定性折扣。"
        "负值如实保留。评分不是命中率，未校准模型最高 C。"
    )
    doc.add_paragraph(
        "官方首发未公布不影响报告生成，也不自动扣分。比赛环境变化仍会改变判断。"
        "本项目不执行真实下注，不提供串关或资金账本。"
    )
