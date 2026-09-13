import json
from dataclasses import FrozenInstanceError, replace
from datetime import timedelta
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document

from football_quant.acquisition.package import load_research
from football_quant.decisions.analysis import analyze_match
from football_quant.domain import (
    Advantage,
    Capability,
    Decision,
    Grade,
    Market,
    Qualitative,
    Report,
    Status,
)
from football_quant.evidence.verification import Claim
from football_quant.models.goals import score_matrix
from football_quant.reporting.word import match_decision_label, needs_deep_analysis, summary, write_report


def research():
    return load_research(Path("fixtures/research-test/manifest.json"))


@pytest.mark.parametrize(
    "market,selection",
    [
        (Market.RESULT, "home"),
        (Market.BTTS, "yes"),
        (Market.GOALS, "2"),
    ],
)
def test_single_quote_ev_without_devig(market, selection) -> None:
    r = research()
    m = r.matches[0]
    q = replace(m.quotes[0], market=market, selection=selection)
    result = analyze_match(replace(m, quotes=(q,)), r)
    c = result.candidates[0]
    assert c.price is not None
    s = score_matrix(result.lambda_home, result.lambda_away)
    if market is Market.RESULT:
        p = sum(v for i, row in enumerate(s.matrix) for j, v in enumerate(row) if i > j)
    elif market is Market.BTTS:
        p = sum(v for i, row in enumerate(s.matrix) for j, v in enumerate(row) if i > 0 and j > 0)
    else:
        p = sum(v for i, row in enumerate(s.matrix) for j, v in enumerate(row) if i + j == 2)
    assert float(c.price.ev) == pytest.approx(p * float(q.decimal_odds) - 1)
    assert c.price.devig_probability is None
    assert c.price.implied_probability == pytest.approx(1 / float(q.decimal_odds))
    assert c.capability is Capability.PARTIAL
    assert all("无法去水比较" not in reason for reason in c.reasons)


def test_goal_baseline_missing_xg_and_lineup() -> None:
    r = research()
    result = analyze_match(r.matches[0], r)
    assert result.lambda_home > 0
    assert result.model_source == "进失球基线（未实证校准）"
    assert "官方首发" not in "；".join(reason for c in result.candidates for reason in c.reasons)
    assert any(c.grade is Grade.C for c in result.candidates)


def test_missing_all_history_has_no_lambda_or_probability() -> None:
    r = research()
    result = analyze_match(replace(r.matches[0], data="{}"), r)
    assert result.lambda_home is result.lambda_away is None
    assert result.scores == ()
    assert all(c.price is None for c in result.candidates)
    assert all(c.confidence is None and c.score_components == () for c in result.candidates)


def test_missing_auxiliary_leaves_main_market_unchanged() -> None:
    r = research()
    m = r.matches[0]
    original = analyze_match(m, r)
    data = json.loads(m.data)
    del data["corners"]
    del data["cards"]
    result = analyze_match(replace(m, data=json.dumps(data)), r)
    assert result.candidates[:3] == original.candidates[:3]
    assert all(c.price is None for c in result.candidates[3:])


@pytest.mark.parametrize(
    "change",
    [
        {"rules": "extra_time"},
        {"original_format": "unknown"},
        {"original_value": "NaN"},
        {"decimal_odds": Decimal("3")},
        {"market": Market.HANDICAP, "line": Decimal("0.3")},
    ],
)
def test_invalid_quote_is_not_priced(change) -> None:
    r = research()
    m = r.matches[0]
    result = analyze_match(replace(m, quotes=(replace(m.quotes[0], **change),)), r)
    assert result.candidates[0].price is None
    assert result.candidates[0].grade is Grade.PASS
    assert result.candidates[0].reasons


def test_stale_or_unknown_odds_time_never_uses_article_time() -> None:
    r = research()
    for observed in (None, r.started - timedelta(hours=7)):
        source = replace(r.evidence[0], observed_at_utc=observed, published_at_utc=r.started)
        result = analyze_match(r.matches[0], replace(r, evidence=(source,)))
        assert all(c.price is None for c in result.candidates)


def test_identity_conflict_blocks_calculation_and_direction() -> None:
    r = research()
    m = r.matches[0]
    claims = m.claims + (Claim("kickoff", "A", "e1"), Claim("kickoff", "B", "e1"))
    result = analyze_match(replace(m, claims=claims), r)
    assert result.data_status is Status.CONFLICT
    assert result.lambda_home is None
    assert all(c.price is None for c in result.candidates)
    assert result.qualitative is None


def test_auxiliary_quote_conflict_is_local() -> None:
    r = research()
    m = r.matches[0]
    q = next(q for q in m.quotes if q.market is Market.CORNER_TOTAL)
    conflicting = replace(q, decimal_odds=Decimal("2.50"), original_value="1.50")
    result = analyze_match(replace(m, quotes=m.quotes + (conflicting,)), r)
    assert result.conflicts
    assert result.candidates[:3] == analyze_match(m, r).candidates[:3]
    assert all(
        c.price is None
        for c in result.candidates
        if c.quote.selection == q.selection and c.quote.market is q.market
    )


def test_negative_ev_preserved_with_decision_separate() -> None:
    r = research()
    m = r.matches[0]
    q = replace(m.quotes[0], decimal_odds=Decimal("1.01"), original_value="1.01")
    c = analyze_match(replace(m, quotes=(q,)), r).candidates[0]
    assert c.price.ev < 0
    assert c.price.risk_ev == c.price.ev - Decimal("0.03")
    assert c.advantage is Advantage.NONPOSITIVE
    assert c.decision is not Decision.UNAVAILABLE
    assert not hasattr(c, "stake")


def test_evidenced_qualitative_without_odds_roundtrips_word(tmp_path) -> None:
    r = research()
    m = r.matches[0]
    qualitative = Qualitative(
        "测试定性主队方向", ("支持证据 [e1]",), ("反向证据 [e1]",), ("输入仍不足",), ("e1",)
    )
    m = replace(m, data="{}", quotes=(), qualitative=qualitative)
    result = analyze_match(m, r)
    assert result.capability is Capability.QUALITATIVE
    assert result.qualitative == qualitative
    assert result.candidates == ()
    assert result.lambda_home is None
    assert result == analyze_match(m, r)
    with pytest.raises(FrozenInstanceError):
        result.qualitative.direction = "changed"
    report = Report(r.mode, r.started, r.deadline, r.generated, r.evidence, (result,), r.coverage)
    path = tmp_path / "qualitative.docx"
    write_report(report, path)
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "方向判断" in text and "测试定性主队方向" in text and "反向证据" in text
    assert "可计算优势" in text and "不产生概率或EV" in text
    assert "资金决策" in text and "无模拟下注金额" in text
    assert "无量化等级" in text
    assert any(rel.target_ref == r.evidence[0].source_url for rel in doc.part.rels.values())


def test_report_separates_qualitative_direction_from_true_pass(tmp_path) -> None:
    r = research()
    first = analyze_match(
        replace(
            r.matches[0],
            data="{}",
            quotes=(),
            qualitative=Qualitative(
                "主队不败方向", ("赛程证据 [e1]",), ("样本不足",), ("不能量化",), ("e1",)
            ),
        ),
        r,
    )
    second = analyze_match(
        replace(r.matches[0], fixture=replace(r.matches[0].fixture, id="m2"), data="{}", quotes=()),
        r,
    )
    report = Report(
        r.mode, r.started, r.deadline, r.generated, r.evidence, (first, second), r.coverage
    )

    assert match_decision_label(first) == "方向观察"
    assert match_decision_label(second) == "PASS"
    assert "定性方向 1 场；PASS 1 场" in summary(report)

    path = tmp_path / "separated.docx"
    write_report(report, path)
    doc = Document(path)
    scan = next(t for t in doc.tables if t.cell(0, 0).text == "北京时间")
    assert [scan.cell(row, 3).text for row in range(1, len(scan.rows))] == ["方向观察", "PASS"]


def test_scan_only_pass_is_not_expanded_as_deep_analysis(tmp_path) -> None:
    r = research()
    direction = analyze_match(
        replace(
            r.matches[0],
            data="{}",
            quotes=(),
            qualitative=Qualitative(
                "主队方向", ("证据 [e1]",), ("反向证据 [e1]",), ("不能量化",), ("e1",)
            ),
        ),
        r,
    )
    scan_only = analyze_match(
        replace(
            r.matches[0],
            fixture=replace(r.matches[0].fixture, id="m2", home="扫描主队", away="扫描客队"),
            data="{}",
            quotes=(),
        ),
        r,
    )
    assert needs_deep_analysis(direction)
    assert not needs_deep_analysis(scan_only)
    report = Report(
        r.mode,
        r.started,
        r.deadline,
        r.generated,
        r.evidence,
        (direction, scan_only),
        r.coverage,
    )
    path = tmp_path / "compact.docx"
    write_report(report, path)
    doc = Document(path)
    deep_headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert f"{direction.fixture.home} 对 {direction.fixture.away}" in deep_headings
    assert "扫描主队 对 扫描客队" not in deep_headings


def test_sample_floor_still_enforced() -> None:
    r = research()
    m = r.matches[0]
    data = json.loads(m.data)
    data["home"] = data["home"][:7]
    result = analyze_match(replace(m, data=json.dumps(data), quotes=m.quotes[:1]), r)
    assert result.lambda_home is None
    assert any("至少8场" in reason for reason in result.missing)


def test_import_rejects_bad_aux_quote_without_losing_main(tmp_path) -> None:
    import shutil

    root = Path("fixtures/research-test")
    raw = json.loads((root / "manifest.json").read_text())
    raw["matches"][0]["quotes"][3]["original_format"] = "unknown"
    shutil.copy(root / "snapshot.txt", tmp_path / "snapshot.txt")
    (tmp_path / "manifest.json").write_text(json.dumps(raw))
    r = load_research(tmp_path / "manifest.json")
    result = analyze_match(r.matches[0], r)
    assert len(result.candidates) == 6
    assert all(c.price for c in result.candidates[:3])
    assert any("拒绝无效报价" in reason for reason in result.missing)


def test_referee_conflict_does_not_reject_goal_model() -> None:
    r = research()
    m = r.matches[0]
    m = replace(m, claims=m.claims + (Claim("referee", "A", "e1"), Claim("referee", "B", "e1")))
    result = analyze_match(m, r)
    assert result.lambda_home > 0
    assert all(c.price for c in result.candidates[:3])
    assert all(c.price is None for c in result.candidates if c.quote.market is Market.CARD_TOTAL)
    assert any("referee" in conflict for conflict in result.conflicts)


def test_xg_path_remains_existing_blend() -> None:
    r = research()
    m = r.matches[0]
    data = json.loads(m.data)
    for side in ("home", "away"):
        for row in data[side]:
            row.update(xg_for=row["goals_for"], xg_against=row["goals_against"], xg_provider="test")
    result = analyze_match(replace(m, data=json.dumps(data)), r)
    baseline = analyze_match(m, r)
    assert result.lambda_home == baseline.lambda_home
    assert result.lambda_away == baseline.lambda_away
    assert result.model_source == "进失球/xG混合基线（未实证校准）"
